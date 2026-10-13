# -*- coding: utf-8 -*-
"""Gera um .docx de exemplo e valida as regras ABNT aplicadas pelo ajustador."""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from abnt import adjust, classify


def build_sample():
    doc = Document()

    # Normal mal formatada (recurso proposital)
    style = doc.styles["Normal"]
    style.font.name = "Comic Sans MS"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 2.5
    style.paragraph_format.left_indent = Cm(2.5)

    # Capa (folha solta)
    p = doc.add_paragraph("Título do Trabalho")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Autor: Fulano de Tal")
    doc.add_page_break()

    # Folha de rosto
    doc.add_paragraph("Título do Trabalho")
    doc.add_paragraph("Autor: Fulano de Tal")
    doc.add_page_break()

    # Resumo
    p = doc.add_paragraph("Resumo")
    p.runs[0].bold = True
    doc.add_paragraph("Este trabalho analisa os impactos da tecnologia. " * 30)
    doc.add_page_break()

    # Abstract
    p = doc.add_paragraph("Abstract")
    p.runs[0].bold = True
    doc.add_paragraph("This paper analyzes the impacts of technology. " * 30)
    doc.add_page_break()

    # Sumário
    p = doc.add_paragraph("Sumário")
    p.runs[0].bold = True
    doc.add_paragraph("1 Introdução\n2 Desenvolvimento\n3 Conclusão")
    doc.add_page_break()

    # Introdução (título numerado primário)
    doc.add_heading("Introdução", level=1)
    doc.add_paragraph("A introdução apresenta o tema e os objetivos do trabalho. " * 40)
    doc.add_paragraph("Lorem ipsum dolor sit amet, consectetur adipiscing elit. " * 50)

    # Citação longa
    quote = "Segundo Silva (2020), a metodologia científica exige rigor. " * 30
    doc.add_paragraph(quote)

    # Desenvolvimento
    doc.add_heading("Desenvolvimento", level=1)
    doc.add_paragraph("O desenvolvimento descreve o método aplicado. " * 40)
    doc.add_heading("Metodologia", level=2)
    doc.add_paragraph("A metodologia foi analítica e descritiva. " * 30)

    # Figura + fonte
    doc.add_paragraph("Figura 1 – Diagrama de fluxo")
    doc.add_paragraph("Fonte: Elaborado pelo autor, 2024.")

    # Tabela com bordas totais (deve perder laterais/verticais)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Coluna A"
    table.cell(0, 1).text = "Coluna B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "4",
            qn("w:space"): "0", qn("w:color"): "auto",
        })
        borders.append(el)
    tbl_pr.append(borders)

    # Conclusão
    doc.add_heading("Conclusão", level=1)
    doc.add_paragraph("Conclui-se que os objetivos foram alcançados. " * 30)

    # Referências
    p = doc.add_paragraph("Referências")
    p.runs[0].bold = True
    doc.add_paragraph("SOBRENOME, Nome. Título da obra. Local: Editora, 2020.")
    doc.add_paragraph("OUTRO, Autor. Outro título. Local: Editora, 2021.")

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample.docx")
    doc.save(out)
    return out


def _header_has_number(header):
    for p in header.paragraphs:
        for instr in p._p.iter(qn("w:instrText")):
            if "PAGE" in (instr.text or ""):
                return True
    return False


def _sect_of_paragraph(doc, paragraph_index):
    for i, p in enumerate(doc.paragraphs):
        sects = p._p.findall(qn("w:pPr") + "/" + qn("w:sectPr"))
        if sects and i <= paragraph_index:
            return sects[0]
    return doc.sections[-1]._sectPr


def _pgnum_attr(sectpr, attr):
    el = sectpr.find(qn("w:pgNumType"))
    if el is None:
        return None
    return el.get(qn(f"w:{attr}"))


def check(result_path):
    doc = Document(result_path)
    errs = []

    def expect(cond, msg):
        if not cond:
            errs.append(msg)
        else:
            print(f"  [OK] {msg}")

    sect = doc.sections[-1]

    # Papel A4
    pw, ph = sect.page_width.cm, sect.page_height.cm
    expect(abs(pw - 21.0) < 0.05 and abs(ph - 29.7) < 0.05, f"A4 ({pw:.2f}x{ph:.2f} cm)")

    # Margens com mirrorMargins (anverso: esq 3, dir 2)
    ml, mr = sect.left_margin.cm, sect.right_margin.cm
    expect(abs(ml - 3.0) < 0.05 and abs(mr - 2.0) < 0.05, f"margens {ml:.2f}/{mr:.2f} cm (3/2)")

    # Mirror margins em settings
    settings = doc.settings.element
    mirror = settings.find(qn("w:mirrorMargins"))
    expect(mirror is not None, "mirrorMargins ativo")

    # Estilo Normal
    style = doc.styles["Normal"]
    expect(
        style.font.name == "Times New Roman",
        f"fonte {style.font.name} = Times New Roman",
    )
    expect(style.font.size == Pt(12), "tamanho Normal = 12")
    expect(
        abs(style.paragraph_format.line_spacing - 1.5) < 0.05,
        "entrelinha Normal = 1,5",
    )

    # Heading 1: sem fonte de tema, 12 pt, preta, negrito
    h1 = doc.styles["Heading 1"]
    h1_rfonts = h1.element.find(qn("w:rPr") + "/" + qn("w:rFonts"))
    expect(
        h1_rfonts is not None and h1_rfonts.get(qn("w:asciiTheme")) is None
        and h1_rfonts.get(qn("w:hAnsiTheme")) is None
        and h1_rfonts.get(qn("w:ascii")) == "Times New Roman",
        "Heading 1 sem tema, Times New Roman",
    )
    expect(h1.font.size == Pt(12) and h1.font.bold, "Heading 1 = 12 negrito")
    expect(
        h1.paragraph_format.space_before == Pt(0)
        and h1.paragraph_format.space_after == Pt(0),
        "Heading 1 sem espaço antes/depois",
    )

    # docDefaults sem tema, TNR 12, 1,5 justificado, recuo 1,25
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    dd_rfonts = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:rFonts"))
    expect(
        dd_rfonts is not None and dd_rfonts.get(qn("w:asciiTheme")) is None
        and dd_rfonts.get(qn("w:ascii")) == "Times New Roman",
        "docDefaults sem tema, Times New Roman",
    )
    dd_spacing = defaults.find(qn("w:pPrDefault") + "/" + qn("w:pPr") + "/" + qn("w:spacing"))
    expect(
        dd_spacing is not None and dd_spacing.get(qn("w:line")) == "360",
        "docDefaults entrelinha 1,5",
    )

    # Mapeamento de títulos
    paras = doc.paragraphs
    heading_idx = {
        i: p for i, p in enumerate(paras) if classify.is_heading(p)
    }
    print(f"  ([INFO] títulos detectados: {sorted(heading_idx)})")

    # 1) RESUMO: centralizado, negrito, MAIÚSCULAS, em nova página
    for i, p in heading_idx.items():
        t = p.text.strip()
        if t.lower() == "resumo":
            fmt = p.paragraph_format
            expect(
                p.alignment == WD_ALIGN_PARAGRAPH.CENTER,
                f"'{t}' centralizado (foi {p.alignment})",
            )
            expect(all(r.bold for r in p.runs), f"'{t}' em negrito")
            expect(fmt.page_break_before, f"'{t}' em nova página")
            expect(t == "RESUMO", f"texto versalete (atual '{t}')")
            break

    # 2) REFERÊNCIAS como título centralizado/uppercase
    for i, p in heading_idx.items():
        t = p.text.strip()
        if t.lower().startswith("refer"):
            expect(
                p.alignment == WD_ALIGN_PARAGRAPH.CENTER,
                f"'{t}' centralizado",
            )
            expect(t == t.upper(), f"'{t}' em MAIÚSCULAS")
            break

    # 3) Títulos numerados: alinhados à esquerda, negrito, nova página
    for i, p in heading_idx.items():
        t = p.text.strip()
        if t.lower() in ("introdução", "desenvolvimento", "conclusão", "metodologia"):
            expect(
                p.alignment == WD_ALIGN_PARAGRAPH.LEFT,
                f"'{t}' alinhado à esquerda",
            )
            expect(all(r.bold for r in p.runs), f"'{t}' negrito")
            level = classify.heading_level(p)
            expect(
                p.paragraph_format.page_break_before or level != 1 or i == 0,
                f"'{t}' page_break antes (nível 1)",
            )

    # 4) Numeração: títulos ≥1 devem ter prefixo e, após o número, um espaço
    #    (o separador entre subníveis pode ser um ponto; o que não pode é
    #    ponto/hífen/travessão logo após o número completo da seção).
    seen_numbered = False
    for i, p in heading_idx.items():
        t = p.text.strip()
        if t and t[0].isdigit():
            seen_numbered = True
            expect(
                re.match(r"^\d+(?:\.\d+)*\s", t) is not None,
                f"numeração com espaço após número em '{t[:40]}'",
            )
            expect(
                re.search(r"\d+[.)-]\s", t) is None,
                f"sem pontuação antes do espaço em '{t[:40]}'",
            )
    expect(seen_numbered, "ao menos um título recebeu número")

    # 5) Resumo corpo: sem recuo, espaçamento simples
    resumo_body = None
    for i, p in enumerate(paras):
        if p.text.strip().lower() == "resumo" and i + 1 < len(paras):
            resumo_body = paras[i + 1]
            break
    if resumo_body is not None:
        fmt = resumo_body.paragraph_format
        fli = fmt.first_line_indent.cm if fmt.first_line_indent else 0.0
        expect(abs(fli - 0.0) < 0.05, "resumo sem recuo de primeira linha")
        expect(
            abs(fmt.line_spacing - 1.0) < 0.05,
            f"resumo espaçamento simples (foi {fmt.line_spacing})",
        )
    else:
        errs.append("resumo: corpo não encontrado")

    # 6) Citação longa: recuo 4 cm, tam 10, espaçamento simples
    quote = None
    for p in paras:
        if p.text.startswith("Segundo Silva"):
            quote = p
            break
    if quote is not None:
        fmt = quote.paragraph_format
        expect(
            abs(fmt.left_indent.cm - 4.0) < 0.05,
            f"citação recuo 4 cm (foi {fmt.left_indent.cm})",
        )
        expect(
            abs(fmt.line_spacing - 1.0) < 0.05,
            "citação espaçamento simples",
        )
        expect(quote.runs and all(r.font.size == Pt(10) for r in quote.runs), "citação tam 10")
    else:
        errs.append("citação longa não detectada")

    # 7) Corpo: recuo 1,25 cm, justificado, entrelinha 1,5
    body = None
    for p in paras:
        if p.text.startswith("A introdução apresenta"):
            body = p
            break
    if body is not None:
        fmt = body.paragraph_format
        expect(
            abs(fmt.first_line_indent.cm - 1.25) < 0.05,
            f"corpo recuo 1,25 cm (foi {fmt.first_line_indent.cm})",
        )
        expect(
            body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY,
            "corpo justificado",
        )
        expect(
            abs(fmt.line_spacing - 1.5) < 0.05,
            "corpo entrelinha 1,5",
        )
    else:
        errs.append("parágrafo de corpo não encontrado")

    # 8) Referência: à esquerda, espaçamento simples, linha em branco
    ref = None
    for i, p in enumerate(paras):
        if p.text.strip().lower().startswith("refer") and i + 1 < len(paras):
            ref = paras[i + 1]
            break
    if ref is not None:
        fmt = ref.paragraph_format
        expect(fmt.alignment == WD_ALIGN_PARAGRAPH.LEFT, "referência à esquerda")
        expect(abs(fmt.line_spacing - 1.0) < 0.05, "referência espaçamento simples")
        sa = fmt.space_after.pt if fmt.space_after else 0.0
        expect(abs(sa - 12.0) < 0.05, f"referência espaço após 12 pt (foi {sa})")

    # 9) Paginação: cabeçalho da última secção tem PAGE em arábico; pré-textual não
    expect(_header_has_number(doc.sections[-1].header), "cabeçalho textual com PAGE")

    # capa: sectPr embutido no 1ª parágrafo deve ter start=0
    para0_sect = doc.paragraphs[0]._p.find(qn("w:pPr") + "/" + qn("w:sectPr"))
    if para0_sect is None:
        para0_sect = doc.paragraphs[0]._p.find(qn("w:pPr"))
        if para0_sect is None:
            expect(False, "capa: sectPr no 1º parágrafo não localizado")
        else:
            para0_sect = para0_sect.find(qn("w:sectPr"))
            if para0_sect is None:
                expect(False, "capa: sectPr no 1º parágrafo não localizado")
            else:
                start = _pgnum_attr(para0_sect, "start")
                expect(start == "0", f"capa start={start} (esperado 0)")
    else:
        start = _pgnum_attr(para0_sect, "start")
        expect(start == "0", f"capa start={start} (esperado 0)")

    # 10) Tabela: sem bordas laterais/verticais
    table = doc.tables[0]
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    for edge in ("left", "right", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        val = (el.get(qn("w:val")) or "").lower() if el is not None else None
        expect(
            el is not None and val == "none",
            f"borda {edge} none (foi {val if el is not None else 'ausente'})",
        )

    print()
    if errs:
        print("FALHAS:")
        for e in errs:
            print(f"  - {e}")
        return False
    print("TODOS OS TESTES PASSARAM")
    return True


def main():
    sample = build_sample()
    print(f"Amostra criada em {sample}")
    result_bytes, report = adjust.adjust_docx(sample, "Times New Roman", "sample.docx")
    result_path = os.path.join(os.path.dirname(sample), "sample_ABNT.docx")
    with open(result_path, "wb") as fh:
        fh.write(result_bytes)
    print(f"Resultado em {result_path}\n")
    print("Relatório de alterações:", report["alteracoes"])
    print("Numeração:", report["numeracao"])
    print()
    ok = check(result_path)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
