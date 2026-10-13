# -*- coding: utf-8 -*-
"""Testa a conversão de um PDF simples em .docx já formatado ABNT."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from abnt import classify, pdf_docx

try:
    import fitz  # PyMuPDF
except ImportError:
    import pymupdf as fitz


def _page(pdf, lines, size=12.0):
    page = pdf.new_page()
    cursor = 30
    for text in lines:
        cursor += size * 1.8
        page.insert_text((30, cursor), text, fontsize=size)
    return page


def build_pdf(path):
    pdf = fitz.open()
    big = 16.0
    _page(pdf, ["Título do Trabalho", "Autor: Fulano de Tal"], size=big)
    _page(pdf, ["Título do Trabalho", "Autor: Fulano de Tal"], size=big)
    _page(pdf, ["Resumo", "Este trabalho analisa os impactos da tecnologia. " * 30], size=big)
    _page(pdf, ["Introdução", "A introdução apresenta o tema do trabalho. " * 40], size=big)
    _page(pdf, ["Referências", "SOBRENOME, Nome. Título. Local: Editora, 2020."], size=big)
    pdf.save(path)
    pdf.close()


def _header_has_number(header):
    for p in header.paragraphs:
        for instr in p._p.iter(qn("w:instrText")):
            if "PAGE" in (instr.text or ""):
                return True
    return False


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    src = os.path.join(base, "sample.pdf")
    out = os.path.join(base, "sample_pdf_ABNT.docx")
    build_pdf(src)

    result_bytes, report = pdf_docx.build_from_pdf(src, "Times New Roman", "sample.pdf")
    with open(out, "wb") as fh:
        fh.write(result_bytes)

    doc = Document(out)
    sect = doc.sections[-1]
    settings = doc.settings.element
    errors = []

    def expect(cond, msg):
        print(f"  [{'OK' if cond else 'ERRO'}] {msg}")
        if not cond:
            errors.append(msg)

    expect(abs(sect.left_margin.cm - 3.0) < 0.05 and abs(sect.right_margin.cm - 2.0) < 0.05,
           f"margens espelhadas {sect.left_margin.cm:.2f}/{sect.right_margin.cm:.2f}")
    expect(settings.find(qn("w:mirrorMargins")) is not None, "mirrorMargins ativo")

    resumo = None
    ref = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == "RESUMO":
            resumo = p
            resumo_body = doc.paragraphs[i + 1]
        if t == "REFERÊNCIAS":
            ref = p
            ref_body = doc.paragraphs[i + 1]

    expect(resumo is not None, "título RESUMO localizado")
    if resumo is not None:
        expect(resumo.alignment == 1, f"resumo centralizado (foi {resumo.alignment})")
        expect(all(r.bold for r in resumo.runs), "resumo negrito")
        expect(resumo.paragraph_format.page_break_before, "resumo em nova página")
        expect(resumo_body.paragraph_format.first_line_indent is None or
               abs(resumo_body.paragraph_format.first_line_indent.cm) < 0.01,
               "resumo sem recuo")
        expect(abs(resumo_body.paragraph_format.line_spacing - 1.0) < 0.05,
               f"resumo espaçamento simples (foi {resumo_body.paragraph_format.line_spacing})")

    expect(ref_body.paragraph_format.space_after is not None and
           abs(ref_body.paragraph_format.space_after.pt - 12) < 0.05,
           f"referência espaço após 12 pt (foi {ref_body.paragraph_format.space_after})")

    expect(_header_has_number(doc.sections[-1].header), "cabeçalho com PAGE")
    pgnum = sect._sectPr.find(qn("w:pgNumType"))
    expect(pgnum is not None and pgnum.get(qn("w:start")) == "1",
           "pgNumType start=1 (foi %s)" % (pgnum.get(qn("w:start")) if pgnum is not None else None))

    print()
    if errors:
        print("FALHAS:")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    print("PDF -> DOCX: TODOS OS TESTES PASSARAM")


if __name__ == "__main__":
    main()
